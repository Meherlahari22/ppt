import os
import re
from typing import TypedDict, List, Optional

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from pptx import Presentation
from pptx.util import Inches, Pt
from langchain_core.messages import BaseMessage, HumanMessage
from langgraph.graph import StateGraph, START, END
from langchain_google_genai import ChatGoogleGenerativeAI


# ============================================================
# 1. CONFIGURATION
# ============================================================

API_KEY = os.getenv("GEMINI_API_KEY")

if not API_KEY:
    raise RuntimeError("GEMINI_API_KEY environment variable is not set.")

MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite-preview")
OUTPUT_DIR = "/tmp/generated_ppts" if os.getenv("RENDER") else "generated_ppts"
os.makedirs(OUTPUT_DIR, exist_ok=True)

llm = ChatGoogleGenerativeAI(
    model=MODEL_NAME,
    google_api_key=API_KEY,
    temperature=0.4,
)


# ============================================================
# 2. STATE
# ============================================================

class PPTState(TypedDict, total=False):
    messages: List[BaseMessage]
    topic: str
    slides: int
    audience: str
    purpose: str
    requirements: str
    outline: str
    content: str
    ppt_path: str
    review: str


# ============================================================
# 3. API REQUEST MODEL
# ============================================================

class PPTRequest(BaseModel):
    topic: str = Field(..., min_length=2)
    slides: int = Field(default=10, ge=3, le=30)
    audience: str = "ECE students"
    purpose: str = "Academic seminar"
    requirements: str = (
        "Include introduction, important concepts, applications, "
        "advantages, limitations and conclusion."
    )


# ============================================================
# 4. HELPER
# ============================================================

def get_text(response) -> str:
    content = response.content

    if isinstance(content, str):
        return content

    if isinstance(content, list):
        parts = []

        for item in content:
            if isinstance(item, dict):
                if "text" in item:
                    parts.append(str(item["text"]))
            else:
                parts.append(str(item))

        return "\n".join(parts)

    return str(content)


# ============================================================
# 5. LANGGRAPH NODES
# ============================================================

def requirement_analyzer(state: PPTState):
    prompt = f"""
You are an expert ECE academic presentation planner.

Analyze this PPT requirement.

Topic: {state["topic"]}
Number of slides: {state["slides"]}
Audience: {state["audience"]}
Purpose: {state["purpose"]}
Additional requirements: {state["requirements"]}

Identify:
1. Main subject
2. Academic level
3. Important technical concepts
4. Required technical depth
5. Suitable presentation flow

Return a concise planning analysis.
"""

    response = llm.invoke(prompt)

    return {
        "requirements": get_text(response),
        "messages": [HumanMessage(content=prompt)],
    }


def ppt_planner(state: PPTState):
    prompt = f"""
You are an expert ECE presentation designer.

Create an outline for exactly {state["slides"]} slides.

Topic: {state["topic"]}
Audience: {state["audience"]}
Purpose: {state["purpose"]}
Requirements: {state["requirements"]}

Rules:
- Slide 1 must be the title slide.
- Keep the flow logical.
- Include technically relevant concepts.
- Include formulas where appropriate.
- Include applications.
- Include advantages/limitations where appropriate.
- End with conclusion.
- Avoid repetition.

Return ONLY:

Slide 1: Title
- ...

Slide 2: Introduction
- ...

...
"""

    response = llm.invoke(prompt)

    return {"outline": get_text(response)}


def content_generator(state: PPTState):
    prompt = f"""
You are an expert ECE professor and technical presentation writer.

Generate PPT-friendly content for this presentation.

Topic: {state["topic"]}
Audience: {state["audience"]}
Purpose: {state["purpose"]}
Slide outline:
{state["outline"]}

Additional requirements:
{state["requirements"]}

For EVERY slide use exactly this structure:

SLIDE X
TITLE: ...

BULLETS:
- ...
- ...
- ...

FORMULA:
Write a formula only if relevant; otherwise write NONE.

DIAGRAM:
Describe a suitable diagram if useful; otherwise write NONE.

Rules:
- Keep bullets concise.
- Do not write long paragraphs.
- Use technically correct formulas.
- Use student-friendly language.
- Do not use markdown tables.
- Do not add extra sections outside the slide structure.
"""

    response = llm.invoke(prompt)

    return {"content": get_text(response)}


# ============================================================
# 6. PPT CREATION
# ============================================================

def parse_slides(content: str):
    pattern = r"SLIDE\s+(\d+)(.*?)(?=SLIDE\s+\d+|$)"

    matches = re.findall(
        pattern,
        content,
        flags=re.IGNORECASE | re.DOTALL,
    )

    slides = []

    for number, block in matches:
        title_match = re.search(
            r"TITLE:\s*(.*?)(?:\n|$)",
            block,
            re.IGNORECASE,
        )

        title = (
            title_match.group(1).strip()
            if title_match
            else f"Slide {number}"
        )

        bullet_match = re.search(
            r"BULLETS:(.*?)(?:FORMULA:|DIAGRAM:|$)",
            block,
            flags=re.IGNORECASE | re.DOTALL,
        )

        bullets = []

        if bullet_match:
            for line in bullet_match.group(1).splitlines():
                line = line.strip()

                if line.startswith("-"):
                    text = line[1:].strip()
                    if text:
                        bullets.append(text)

        formula_match = re.search(
            r"FORMULA:\s*(.*?)(?:DIAGRAM:|$)",
            block,
            flags=re.IGNORECASE | re.DOTALL,
        )

        formula = (
            formula_match.group(1).strip()
            if formula_match
            else ""
        )

        if formula.upper() == "NONE":
            formula = ""

        diagram_match = re.search(
            r"DIAGRAM:\s*(.*)",
            block,
            flags=re.IGNORECASE | re.DOTALL,
        )

        diagram = (
            diagram_match.group(1).strip()
            if diagram_match
            else ""
        )

        if diagram.upper() == "NONE":
            diagram = ""

        slides.append(
            {
                "number": number,
                "title": title,
                "bullets": bullets,
                "formula": formula,
                "diagram": diagram,
            }
        )

    return slides


def create_ppt(state: PPTState):
    prs = Presentation()

    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = parse_slides(state["content"])

    if not slides:
        raise ValueError("Could not parse slide content generated by Gemini.")

    for index, slide_data in enumerate(slides):
        # Title layout for first slide
        if index == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])

            slide.shapes.title.text = slide_data["title"]

            if len(slide.placeholders) > 1:
                subtitle = slide.placeholders[1]
                subtitle.text = (
                    f"{state['purpose']}\n"
                    f"{state['audience']}"
                )

        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])

            slide.shapes.title.text = slide_data["title"]

            title_paragraph = slide.shapes.title.text_frame.paragraphs[0]
            title_paragraph.font.size = Pt(28)

            body = slide.placeholders[1]
            body.text = ""

            text_frame = body.text_frame
            text_frame.word_wrap = True

            for bullet in slide_data["bullets"]:
                p = text_frame.add_paragraph()
                p.text = bullet
                p.level = 0
                p.font.size = Pt(19)

            if slide_data["formula"]:
                p = text_frame.add_paragraph()
                p.text = "Formula: " + slide_data["formula"]
                p.font.size = Pt(17)

            if slide_data["diagram"]:
                p = text_frame.add_paragraph()
                p.text = "Diagram: " + slide_data["diagram"]
                p.font.size = Pt(15)

    safe_topic = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        state["topic"],
    ).strip("_")

    if not safe_topic:
        safe_topic = "ECE_Presentation"

    ppt_path = os.path.join(
        OUTPUT_DIR,
        f"{safe_topic}.pptx",
    )

    prs.save(ppt_path)

    return {"ppt_path": ppt_path}


def ppt_reviewer(state: PPTState):
    prompt = f"""
You are a senior ECE professor reviewing a technical presentation.

Topic:
{state["topic"]}

Outline:
{state["outline"]}

Generated content:
{state["content"]}

Check:
1. Technical correctness
2. Logical flow
3. ECE relevance
4. Slide organization
5. Missing concepts
6. Unnecessary repetition
7. Grammar

Give a concise review and finish with:
STATUS: APPROVED
"""

    response = llm.invoke(prompt)

    return {"review": get_text(response)}


# ============================================================
# 7. LANGGRAPH
# ============================================================

workflow = StateGraph(PPTState)

workflow.add_node("requirement_analyzer", requirement_analyzer)
workflow.add_node("ppt_planner", ppt_planner)
workflow.add_node("content_generator", content_generator)
workflow.add_node("ppt_generator", create_ppt)
workflow.add_node("ppt_reviewer", ppt_reviewer)

workflow.add_edge(START, "requirement_analyzer")
workflow.add_edge("requirement_analyzer", "ppt_planner")
workflow.add_edge("ppt_planner", "content_generator")
workflow.add_edge("content_generator", "ppt_generator")
workflow.add_edge("ppt_generator", "ppt_reviewer")
workflow.add_edge("ppt_reviewer", END)

ppt_app = workflow.compile()


# ============================================================
# 8. FASTAPI
# ============================================================

app = FastAPI(
    title="ECE PPT Generation Agent",
    description="LangGraph + Gemini PowerPoint Generation Agent",
    version="1.0.0",
)


@app.get("/")
async def root():
    return {
        "message": "ECE PPT Generation Agent is running.",
        "endpoint": "/agent/playground",
        "method": "POST",
    }


@app.post("/agent/playground")
async def agent_playground(request: PPTRequest):
    try:
        initial_state: PPTState = {
            "messages": [],
            "topic": request.topic,
            "slides": request.slides,
            "audience": request.audience,
            "purpose": request.purpose,
            "requirements": request.requirements,
        }

        result = ppt_app.invoke(
            initial_state,
            config={"recursion_limit": 30},
        )

        ppt_path = result.get("ppt_path")

        if not ppt_path or not os.path.exists(ppt_path):
            raise HTTPException(
                status_code=500,
                detail="PPT generation failed.",
            )

        return {
            "status": "success",
            "message": "PPT generated successfully.",
            "topic": request.topic,
            "slides": request.slides,
            "download_url": "/agent/download",
            "review": result.get("review", ""),
        }

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e),
        )


@app.get("/agent/download")
async def download_ppt():
    if not os.path.exists(OUTPUT_DIR):
        raise HTTPException(
            status_code=404,
            detail="No PPT has been generated.",
        )

    ppt_files = [
        filename
        for filename in os.listdir(OUTPUT_DIR)
        if filename.lower().endswith(".pptx")
    ]

    if not ppt_files:
        raise HTTPException(
            status_code=404,
            detail="No PPT has been generated yet.",
        )

    latest_file = max(
        ppt_files,
        key=lambda filename: os.path.getmtime(
            os.path.join(OUTPUT_DIR, filename)
        ),
    )

    file_path = os.path.join(
        OUTPUT_DIR,
        latest_file,
    )

    return FileResponse(
        file_path,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
        filename=latest_file,
    )


# ============================================================
# 9. LOCAL/RENDER START
# ============================================================

if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        "app:app",
        host="0.0.0.0",
        port=int(os.getenv("PORT", "8000")),
    )
